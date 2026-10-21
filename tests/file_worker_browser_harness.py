"""Serve the actual packaged file worker with synthetic, library-made documents.

Run with the bundled workspace Python, which supplies the document libraries.
This checks browser execution of the packaged parser, not native extension
permissions, upload association, or Chrome/Edge/Firefox release certification.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import mimetypes
import re
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote, urlsplit

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument(
    "--build",
    choices=[
        target + suffix
        for target in ("chrome-mv3", "edge-mv3", "firefox-mv2")
        for suffix in ("", "-original", "-ecologits")
    ],
    default="chrome-mv3",
    help="Read the worker from this existing build; this does not install an extension.",
)
parser.add_argument("--port", type=int, default=4183)
options = parser.parse_args()
if not 1 <= options.port <= 65535:
    parser.error("--port must be between 1 and 65535")
BUILD = ROOT / "apps/extension/.output" / options.build
MARKER = "FERN-LOCAL-042"
HOST = f"file-worker.localhost:{options.port}"
ALLOWED_HOSTS = {HOST, f"127.0.0.1:{options.port}"}
if options.port == 80:
    ALLOWED_HOSTS.update({"file-worker.localhost", "127.0.0.1"})


def cdata_copy(payload: bytes, part: str, mixed: bool) -> bytes:
    """Alter only synthetic XML text while retaining a valid Office ZIP container."""
    output = io.BytesIO()
    replacement = b"FERN-<![CDATA[LOCAL]]>-042" if mixed else b"<![CDATA[FERN-LOCAL-042]]>"
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == part:
                assert data.count(MARKER.encode()) == 1
                data = data.replace(MARKER.encode(), replacement)
            target.writestr(entry, data)
    return output.getvalue()


def fixtures() -> dict[str, tuple[str, bytes, str]]:
    values = {"text": ("text", MARKER.encode(), "counted")}
    document = Document()
    document.core_properties.author = "Synthetic verification"
    document.add_paragraph(MARKER)
    output = io.BytesIO()
    document.save(output)
    values["docx"] = ("docx", output.getvalue(), "counted")

    workbook = Workbook()
    workbook.properties.creator = "Synthetic verification"
    workbook.active["A1"] = MARKER
    workbook.active["A2"] = 42
    output = io.BytesIO()
    workbook.save(output)
    values["xlsx"] = ("xlsx", output.getvalue(), "counted")

    presentation = Presentation()
    presentation.core_properties.author = "Synthetic verification"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = MARKER
    output = io.BytesIO()
    presentation.save(output)
    values["pptx"] = ("pptx", output.getvalue(), "counted")

    for name in ("pdf", "encrypted-pdf", "scanned-pdf"):
        output = io.BytesIO()
        canvas = Canvas(output, encrypt="synthetic-password" if name == "encrypted-pdf" else None)
        canvas.setAuthor("Synthetic verification")
        if name == "scanned-pdf":
            canvas.drawImage(ImageReader(Image.new("RGB", (16, 16), "green")), 72, 700)
        else:
            canvas.drawString(72, 720, MARKER)
        canvas.save()
        expected = {"pdf": "counted", "encrypted-pdf": "encrypted_file", "scanned-pdf": "no_extractable_text"}[name]
        values[name] = ("pdf", output.getvalue(), expected)

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Hello")
    hidden = paragraph.add_run()
    hidden.font.hidden = True
    hidden.add_tab()
    hidden.add_break()
    paragraph.add_run("Hello")
    output = io.BytesIO()
    document.save(output)
    values["docx-hidden-separators"] = ("docx", output.getvalue(), "counted")
    for kind, part in {
        "docx": "word/document.xml",
        "xlsx": "xl/worksheets/sheet1.xml",
        "pptx": "ppt/slides/slide1.xml",
    }.items():
        for mixed in (False, True):
            name = kind + ("-cdata-mixed" if mixed else "-cdata-only")
            values[name] = (kind, cdata_copy(values[kind][1], part, mixed), "counted")
    return values


def worker_path() -> str:
    references = set()
    # Firefox references the shared worker from its background rather than an offscreen page.
    for source in BUILD.rglob("*.js"):
        references.update(re.findall(r"/assets/file\.worker-[A-Za-z0-9_-]+\.js", source.read_text(encoding="utf-8")))
    if len(references) != 1:
        raise RuntimeError("Expected one file worker referenced by the selected build")
    return references.pop()


FIXTURES = fixtures()
WORKER = worker_path()
WORKER_HASH = hashlib.sha256((BUILD / WORKER.lstrip("/")).read_bytes()).hexdigest()
# Supplied-text oracles independently checked with official tiktoken 0.14.0,
# encode_ordinary. These concern extracted text, not provider request usage.
TOKEN_ORACLES = {"text": "6", "pdf": "7", "docx": "7", "xlsx": "9", "pptx": "7"}
TOKEN_ORACLES["docx-hidden-separators"] = "3"
for _kind in ("docx", "xlsx", "pptx"):
    for _suffix in ("-cdata-only", "-cdata-mixed"):
        TOKEN_ORACLES[_kind + _suffix] = TOKEN_ORACLES[_kind]
CASES = [{"name": name, "kind": value[0], "expected": value[2], "tokens": TOKEN_ORACLES.get(name)} for name, value in FIXTURES.items()]
PAGE = r'''<!doctype html>
<html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>Packaged file worker verification</title>
<style>body{font:16px system-ui;max-width:75rem;margin:2rem auto;padding:1rem;color:#203524;background:#f7f9f2}button{font:inherit;padding:.7rem 1rem}pre{white-space:pre-wrap;overflow-wrap:anywhere;border:1px solid #8a9c83;padding:1rem}strong{display:block;margin:1rem 0}</style>
<h1>Packaged file worker verification</h1>
<p>__COUNT__ synthetic documents are parsed locally using the built extension worker. No provider is contacted. Native extension permissions and upload association are outside this check.</p>
<p>Build: <code>__BUILD__</code>. Worker SHA256: <code>__WORKER_HASH__</code>. This page tests the selected worker in the browser displaying this page, not every target browser.</p>
<button id="run">Run local extraction checks</button><strong id="status" role="status">Ready</strong><div id="results"></div>
<script type="module">
const cases = __CASES__, workerUrl = __WORKER__;
const status = document.querySelector('#status'), results = document.querySelector('#results'), button = document.querySelector('#run');
function count(job) {
  return new Promise(resolve => {
    let worker;
    const finish = result => { clearTimeout(timer); worker?.terminate(); resolve(result); };
    const timer = setTimeout(() => finish({status:'harness-failure',reason:'timeout'}), 32000);
    try {
      worker = new Worker(workerUrl, {type:'module'});
      worker.onmessage = event => finish(event.data);
      worker.onerror = () => finish({status:'harness-failure',reason:'worker_error'});
      worker.postMessage(job, [job.bytes]);
    } catch { finish({status:'harness-failure',reason:'worker_unavailable'}); }
  });
}
button.addEventListener('click', async () => {
  button.disabled = true;
  results.replaceChildren();
  let passed = 0;
  try {
    for (const item of cases) {
      status.textContent = 'Checking ' + item.name;
      const response = await fetch('/fixture/' + item.name);
      if (!response.ok) throw new Error('fixture_unavailable');
      const bytes = await response.arrayBuffer(), size = bytes.byteLength;
      const result = await count({jobId:crypto.randomUUID(),revision:1,kind:item.kind,bytes});
      const numeric = item.expected === 'counted' ? result.status === 'counted' && result.tokens === item.tokens : result.status === 'unavailable' && result.tokens === null && result.reason === item.expected;
      const correct = numeric && result.inputBytes === size && result.scope === 'extracted_file_text' && result.association === 'unassociated' && result.telemetryAllowed === false && !JSON.stringify(result).includes('FERN-LOCAL-042');
      if (correct) passed++;
      const pre = document.createElement('pre');
      pre.textContent = (correct ? 'PASS ' : 'FAIL ') + item.name + '\n' + JSON.stringify(result, null, 2);
      results.append(pre);
    }
    status.textContent = passed + ' of ' + cases.length + ' checks passed';
  } catch (error) { status.textContent = 'Harness failed: ' + error.message; }
  finally { button.disabled = false; }
});
</script></html>'''.replace("__COUNT__", str(len(CASES))).replace("__CASES__", json.dumps(CASES)).replace("__WORKER__", json.dumps(WORKER)).replace("__BUILD__", options.build).replace("__WORKER_HASH__", WORKER_HASH).encode()


class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.headers.get("Host") not in ALLOWED_HOSTS:
            self.send_error(403)
            return
        path = unquote(urlsplit(self.path).path)
        if path == "/":
            body, mime = PAGE, "text/html; charset=utf-8"
        elif path.startswith("/fixture/") and path[9:] in FIXTURES:
            body, mime = FIXTURES[path[9:]][1], "application/octet-stream"
        elif path.startswith("/assets/"):
            asset = (BUILD / path.lstrip("/")).resolve()
            if not asset.is_relative_to((BUILD / "assets").resolve()) or not asset.is_file():
                self.send_error(404)
                return
            body = asset.read_bytes()
            mime = "text/javascript" if asset.suffix in {".js", ".mjs"} else mimetypes.guess_type(asset.name)[0] or "application/octet-stream"
        else:
            self.send_error(404)
            return
        self.send_response(200)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'unsafe-inline'; worker-src 'self'; connect-src 'self'")
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *_):
        pass


if __name__ == "__main__":
    print(f"Synthetic packaged-worker harness: http://{HOST}/", flush=True)
    print(f"Selected build: {options.build}", flush=True)
    print(f"Referenced worker: {WORKER}", flush=True)
    print(f"Worker SHA256: {WORKER_HASH}", flush=True)
    ThreadingHTTPServer(("127.0.0.1", options.port), Handler).serve_forever()
